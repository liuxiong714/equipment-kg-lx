"""DOCX 解析器 — python-docx。"""
from __future__ import annotations

import io

from parsers.base import BaseParser, register_parser


@register_parser(".docx")
class DocxParser(BaseParser):
    """Word 文档提取。"""

    supported_ext = ".docx"

    def extract(self, file_bytes: bytes) -> str:
        try:
            from docx import Document
        except ImportError:
            raise RuntimeError("python-docx 未安装。请: py -3.12 -m pip install python-docx")

        doc = Document(io.BytesIO(file_bytes))
        paras = [p.text for p in doc.paragraphs if p.text.strip()]
        # 表格
        for tbl in doc.tables:
            for row in tbl.rows:
                paras.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n".join(paras)
